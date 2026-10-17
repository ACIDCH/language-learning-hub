import io
import json
import zipfile
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
from server.app import create_app, extract_document

@pytest.fixture
def client(tmp_path):
    course = {'id':'fr-a1-u01','title':'Bonjour','language':'fr','level':'A1','unit':1,'text':'Original preserved', 'sections':[], 'exercises':[{'id':'closed','mode':'exact_token','answer':'êtes','explanation':'vous form'},{'id':'open','mode':'self_review','answer':'An example','explanation':'Many answers accepted'}], 'audio':[], 'versions':[]}
    (tmp_path/'courses.json').write_text(json.dumps([course]),encoding='utf-8')
    return TestClient(create_app(tmp_path))

def test_catalog_detail_and_missing(client):
    assert client.get('/api/courses').json()[0]['id']=='fr-a1-u01'
    assert client.get('/api/courses/fr-a1-u01').json()['text']=='Original preserved'
    assert client.get('/api/courses/missing').status_code==404

def test_closed_gap_correction_keeps_accents_and_open_ended_review(client):
    for value,result in [(' êtes ','correct'),('e\u0302tes','correct'),('etes','review')]:
        assert client.post('/api/courses/fr-a1-u01/check',json={'exercise_id':'closed','answer':value}).json()['result']==result
    assert client.post('/api/courses/fr-a1-u01/check',json={'exercise_id':'open','answer':'A valid alternative'}).json()['result']=='self_review'

def test_progress_survives_restart(client):
    state={'stage':4,'notes':'épreuve 中文','favorite':True,'answers':{'q1':{'text':'Bonjour','result':'review'}}}
    assert client.put('/api/state/fr-a1-u01',json=state).status_code==200
    value=client.get('/api/state').json()['fr-a1-u01']
    assert value['stage']==4 and value['answers']['q1']['result']=='review'
    assert client.put('/api/state/fr-a1-u01',json={'stage':7}).status_code==422
    assert client.put('/api/state/missing',json=state).status_code==404

def test_revisions_are_append_only(client):
    for text in ['First correction','Second correction']:
        assert client.post('/api/courses/fr-a1-u01/revisions',json={'text':text}).status_code==200
    value=client.get('/api/courses/fr-a1-u01').json()
    assert value['text']=='Original preserved'
    assert len(value['personal_revisions'])==2

def test_cloud_restore_keeps_edit_time_and_rejects_older_copy(client):
    url = '/api/state/fr-a1-u01/restore'
    cloud = {'state': {'stage': 3, 'notes': 'Cloud épreuve', 'answers': {'q1': {'text': 'êtes', 'result': 'correct'}}}, 'updated_at': '2026-01-02T03:04:05+13:00'}
    result = client.post(url, json=cloud)
    assert result.status_code == 200
    restored = client.get('/api/state').json()['fr-a1-u01']
    assert restored['updated_at'] == '2026-01-01T14:04:05+00:00'
    assert restored['answers']['q1']['result'] == 'correct'
    assert client.post(url, json=cloud).json() == restored
    older = {'state': {'notes': 'Stale cloud copy'}, 'updated_at': '2026-01-01T00:00:00Z'}
    assert client.post(url, json=older).json() == restored
    client.put('/api/state/fr-a1-u01', json={'notes': 'New local edit'})
    assert client.post(url, json=cloud).json()['notes'] == 'New local edit'
    assert client.post(url, json={'state': {}, 'updated_at': '2026-01-01T00:00:00'}).status_code == 422
    assert client.post('/api/state/missing/restore', json=cloud).status_code == 404

def test_import_and_dedupe(client):
    raw='# Chapitre 1\nBonjour à tous.\n## Chapitre 2\nFin.'.encode()
    first=client.post('/api/library',files={'file':('../../sample.md',raw,'text/markdown')}).json()
    assert not first['duplicate']
    item=client.get('/api/library/'+first['id']).json()
    assert item['name']=='sample.md' and item['text']==raw.decode()
    assert client.get('/api/library/'+first['id']+'/original').content==raw
    assert client.post('/api/library',files={'file':('copy.txt',raw)}).json()['duplicate']

def test_docx_order():
    d=Document();d.add_paragraph('Before');d.add_table(rows=1,cols=1).cell(0,0).text='Inside';d.add_paragraph('After')
    stream=io.BytesIO();d.save(stream)
    assert extract_document(stream.getvalue(),'.docx')=='Before\nInside\nAfter'

def test_epub_spine_order():
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w') as z:
        z.writestr('META-INF/container.xml','<container><rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>')
        z.writestr('OPS/book.opf','<package xmlns="http://www.idpf.org/2007/opf"><manifest><item id="a" href="a.xhtml"/><item id="b" href="b.xhtml"/></manifest><spine><itemref idref="b"/><itemref idref="a"/></spine></package>')
        z.writestr('OPS/a.xhtml','<p>Second</p>');z.writestr('OPS/b.xhtml','<p>First</p>')
    assert extract_document(stream.getvalue(),'.epub')=='First\n\nSecond'

def test_pdf_text_and_scanned_rejection():
    w=PdfWriter();p=w.add_blank_page(width=300,height=300)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    p[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    content=DecodedStreamObject();content.set_data(b'BT /F1 12 Tf 20 200 Td (Bonjour) Tj ET');p[NameObject('/Contents')]=content
    b=io.BytesIO();w.write(b);assert 'Bonjour' in extract_document(b.getvalue(),'.pdf')
    w=PdfWriter();w.add_blank_page(width=300,height=300);b=io.BytesIO();w.write(b)
    with pytest.raises(ValueError,match='OCR'):extract_document(b.getvalue(),'.pdf')

def test_bad_imports_and_cross_site(client):
    assert client.post('/api/library',files={'file':('fake.mp3',b'not audio')}).status_code==422
    assert client.post('/api/library',files={'file':('video.mp4',b'video')}).status_code==422
    assert client.get('/api/audio/unknown.mp3').status_code==404
    assert client.put('/api/state/fr-a1-u01',json={},headers={'Origin':'https://evil.example'}).status_code==403
    assert client.get('/api/state',headers={'Host':'evil.example'}).status_code==400
    assert client.post('/api/ai',json={'provider':'bad','prompt':'test'}).status_code==422

def test_audio_range_and_integrity(tmp_path):
    directory=tmp_path/'audio';directory.mkdir()
    (directory/'manifest.json').write_text(json.dumps([{'filename':'sample.mp3'}]))
    (directory/'sample.mp3').write_bytes(b'0123456789')
    c=TestClient(create_app(tmp_path));r=c.get('/api/audio/sample.mp3',headers={'Range':'bytes=2-5'})
    assert r.status_code==206 and r.content==b'2345'

def test_data_must_stay_outside_repo():
    with pytest.raises(RuntimeError,match='outside'):create_app(Path(__file__).resolve().parent)

def test_ai_settings_never_return_key(client):
    key='test-only-not-a-real-credential'
    assert client.put('/api/ai/settings',json={'remote_key':key,'remote_model':'test-model'}).status_code==200
    result=client.get('/api/ai/settings')
    assert result.json()['key_configured'] and key not in result.text
    assert client.put('/api/ai/settings',json={'remote_key':None,'remote_model':'next-model'}).status_code==200
    assert client.get('/api/ai/settings').json()['key_configured']
    assert client.put('/api/ai/settings',json={'remote_key':''}).status_code==200
    assert not client.get('/api/ai/settings').json()['key_configured']
    assert client.put('/api/ai/settings',json={'remote_url':'http://example.test'}).status_code==422
